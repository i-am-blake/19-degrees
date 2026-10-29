"""Convert the Australian Shark-Incident Database spreadsheet (Zenodo) to the CSV the build scripts read.
Usage: python3 -I asid_to_csv.py "<Australian Shark-Incident Database Public Version.xlsx>" asid.csv
Download the newest version from https://zenodo.org/records/18752301 (CC BY 4.0)."""
import csv, sys
from openpyxl import load_workbook
src, out = sys.argv[1:3]
ws = load_workbook(src, read_only=True, data_only=True).worksheets[0]
rows = ws.iter_rows(values_only=True)
head = next(rows)
keep = [i for i, h in enumerate(head) if h not in (None, "")]
yi = head.index("Incident.year")
with open(out, "w", newline="", encoding="utf-8") as f:
    w = csv.writer(f)
    w.writerow([head[i] for i in keep])
    n = 0
    for r in rows:
        if r[yi] in (None, ""): continue
        w.writerow(["" if r[i] is None else r[i] for i in keep]); n += 1
print(f"wrote {n} incidents to {out}")
